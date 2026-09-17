"""Tests de los endpoints de ``apps.fe.views.certificacion_paso2_ecf_view``
/ ``certificacion_paso2_rfce_view`` / ``certificacion_paso3_view`` --
version "bulk" (sube el Excel oficial de la DGII completo) de lo que
``pruebas_enviar_view`` hace fila por fila.

Misma estrategia que ``test_views_pruebas.py``: nunca se toca Oracle real
ni la DGII real -- se monkeypatchea ``ecf_builder``/``dgii_client`` en la
frontera.
"""
from __future__ import annotations

import io

import openpyxl
import pytest
from django.contrib.auth import get_user_model
from django.test import Client

from apps.fe import dgii_client, ecf_builder


@pytest.fixture
def cliente_autenticado(db):
    User = get_user_model()
    user = User.objects.create_user(username='tester', password='x')
    client = Client()
    client.force_login(user)
    return client


def _excel_ecf_bytes(filas):
    """Arma un .xlsx en memoria con hoja 'ECF', misma forma que el Excel
    real de la DGII: fila 0 = headers, resto = datos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ECF'
    headers = ['TipoeCF', 'ENCF', 'RNCEmisor', 'RNCComprador', 'MontoTotal']
    ws.append(headers)
    for fila in filas:
        ws.append([fila.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def test_requiere_login(client, db):
    resp = client.post('/api/fe/certificacion/paso2-ecf/', data={'no_cia': '01'})
    assert resp.status_code in (302, 401, 403)


def test_sin_archivo_da_400(cliente_autenticado):
    resp = cliente_autenticado.post(
        '/api/fe/certificacion/paso2-ecf/', data={'no_cia': '01'})
    assert resp.status_code == 400
    assert 'archivo' in resp.json()['detail'].lower()


def test_envia_cada_fila_y_reporta_resultados(cliente_autenticado, monkeypatch):
    filas = [
        {'TipoeCF': 32, 'ENCF': 'E320000000006', 'RNCEmisor': '130217432',
         'RNCComprador': '131880681', 'MontoTotal': 1180},
        {'TipoeCF': 31, 'ENCF': 'E310000000001', 'RNCEmisor': '130217432',
         'RNCComprador': '131880681', 'MontoTotal': 7080},
    ]
    archivo = _excel_ecf_bytes(filas)

    build_calls = []
    monkeypatch.setattr(
        ecf_builder, 'construir_ecf_generico',
        lambda tipo, encf, datos: build_calls.append((tipo, encf)) or '<ECF/>')

    envios = []

    def fake_enviar_ecf(no_cia, ambiente, e_ncf, xml_sin_firmar):
        envios.append((ambiente, e_ncf))
        return {'trackId': f'TRACK-{e_ncf}', 'xml_firmado': '<ECF firmado/>',
                'respuesta_cruda': {'trackId': f'TRACK-{e_ncf}'}}

    monkeypatch.setattr(dgii_client, 'enviar_ecf', fake_enviar_ecf)
    monkeypatch.setattr(
        'apps.legacy.repositories.fe_repo.save_documento_enviado',
        lambda *a, **k: None)

    resp = cliente_autenticado.post(
        '/api/fe/certificacion/paso2-ecf/',
        data={'no_cia': '01', 'archivo': _as_upload(archivo)},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert len(body['resultados']) == 2
    assert all(r['ok'] for r in body['resultados'])
    assert {e for _, e in envios} == {'E320000000006', 'E310000000001'}
    # SIEMPRE certecf, sin importar TFE_CONFIG.ambiente -- mismo requisito
    # de seguridad que pruebas_enviar_view.
    assert all(amb == 'certecf' for amb, _ in envios)


def test_fila_con_error_no_detiene_las_demas(cliente_autenticado, monkeypatch):
    filas = [
        {'TipoeCF': 32, 'ENCF': 'E320000000006', 'RNCEmisor': '130217432',
         'RNCComprador': '131880681', 'MontoTotal': 1180},
        {'TipoeCF': 31, 'ENCF': 'E310000000001', 'RNCEmisor': '130217432',
         'RNCComprador': '131880681', 'MontoTotal': 7080},
    ]
    archivo = _excel_ecf_bytes(filas)

    def fake_construir(tipo, encf, datos):
        if encf == 'E320000000006':
            raise ecf_builder.ECFBuilderError('dato invalido')
        return '<ECF/>'

    monkeypatch.setattr(ecf_builder, 'construir_ecf_generico', fake_construir)
    monkeypatch.setattr(
        dgii_client, 'enviar_ecf',
        lambda no_cia, ambiente, e_ncf, xml: {
            'trackId': 'T1', 'xml_firmado': '<x/>', 'respuesta_cruda': {}})
    monkeypatch.setattr(
        'apps.legacy.repositories.fe_repo.save_documento_enviado',
        lambda *a, **k: None)

    resp = cliente_autenticado.post(
        '/api/fe/certificacion/paso2-ecf/',
        data={'no_cia': '01', 'archivo': _as_upload(archivo)},
    )
    assert resp.status_code == 200
    resultados = {r['encf']: r for r in resp.json()['resultados']}
    assert resultados['E320000000006']['ok'] is False
    assert resultados['E310000000001']['ok'] is True


def _as_upload(contenido: bytes):
    from django.core.files.uploadedfile import SimpleUploadedFile
    return SimpleUploadedFile(
        'set-pruebas.xlsx', contenido,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _excel_rfce_bytes(filas_ecf, filas_rfce):
    wb = openpyxl.Workbook()
    ws_ecf = wb.active
    ws_ecf.title = 'ECF'
    headers = ['TipoeCF', 'ENCF', 'RNCEmisor', 'RNCComprador', 'MontoTotal']
    ws_ecf.append(headers)
    for fila in filas_ecf:
        ws_ecf.append([fila.get(h) for h in headers])
    ws_rfce = wb.create_sheet('RFCE')
    ws_rfce.append(headers)
    for fila in filas_rfce:
        ws_rfce.append([fila.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def test_paso2_rfce_requiere_login(client, db):
    resp = client.post('/api/fe/certificacion/paso2-rfce/', data={'no_cia': '01'})
    assert resp.status_code in (302, 401, 403)


def test_paso2_rfce_devuelve_xml_firmado_para_descargar(cliente_autenticado, monkeypatch):
    fila = {'TipoeCF': 32, 'ENCF': 'E320000000012', 'RNCEmisor': '130217432',
            'RNCComprador': '131880681', 'MontoTotal': 47200}
    archivo = _excel_rfce_bytes([fila], [fila])

    monkeypatch.setattr(ecf_builder, 'construir_ecf_generico', lambda t, e, d: '<ECF/>')
    monkeypatch.setattr(
        dgii_client, '_firmar_para_envio',
        lambda no_cia, xml: ('<ECF firmado><Signature><SignatureValue>abc123XYZ==</SignatureValue></Signature></ECF>', '130217432'))
    monkeypatch.setattr(
        ecf_builder, 'derivar_codigo_seguridad', lambda xml_firmado: 'abc123')
    monkeypatch.setattr(ecf_builder, 'construir_rfce', lambda e, d, c: '<RFCE/>')
    monkeypatch.setattr(
        dgii_client, 'enviar_rfce',
        lambda no_cia, ambiente, e_ncf, xml: {
            'estado': 'Aceptado', 'codigo': 1, 'mensajes': None,
            'encf': e_ncf, 'secuencia_utilizada': True,
            'xml_firmado': '<RFCE firmado/>', 'respuesta_cruda': {}})

    resp = cliente_autenticado.post(
        '/api/fe/certificacion/paso2-rfce/',
        data={'no_cia': '01', 'archivo': _as_upload(archivo)},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert len(body['resultados']) == 1
    r = body['resultados'][0]
    assert r['ok'] is True
    assert r['encf'] == 'E320000000012'
    assert r['estado_rfce'] == 'Aceptado'
    assert 'SignatureValue' in r['ecf32_firmado_xml']
    assert r['nombre_archivo'] == '130217432E320000000012.xml'


def _excel_aprobaciones_bytes(filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ACEECF_Generadas'
    headers = ['Version', 'RNCEmisor', 'eNCF', 'FechaEmision', 'MontoTotal',
               'RNCComprador', 'Estado', 'DetalleMotivoRechazo',
               'FechaHoraAprobacionComercial']
    ws.append(headers)
    for fila in filas:
        ws.append([fila.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def test_paso3_requiere_login(client, db):
    resp = client.post('/api/fe/certificacion/paso3-aprobaciones/', data={'no_cia': '01'})
    assert resp.status_code in (302, 401, 403)


def test_paso3_envia_cada_fila_del_excel(cliente_autenticado, monkeypatch):
    fila = {
        'Version': '1.0', 'RNCEmisor': '131880681', 'eNCF': 'E310000000001',
        'FechaEmision': '01-04-2020', 'MontoTotal': 7080,
        'RNCComprador': '130217432', 'Estado': 1,
        'DetalleMotivoRechazo': None,
        'FechaHoraAprobacionComercial': '17-09-2026 11:44:23',
    }
    archivo = _excel_aprobaciones_bytes([fila])

    monkeypatch.setattr(ecf_builder, 'construir_acecf', lambda row: '<ACECF/>')
    envios = []

    def fake_enviar(no_cia, ambiente, e_ncf, rnc_comprador, xml):
        envios.append((ambiente, e_ncf, rnc_comprador))
        return {'mensaje': [], 'estado': 'Aprobacion Comercial Aprobada.',
                'codigo': '01', 'xml_firmado': '<x/>', 'respuesta_cruda': {}}

    monkeypatch.setattr(dgii_client, 'enviar_aprobacion_comercial', fake_enviar)

    resp = cliente_autenticado.post(
        '/api/fe/certificacion/paso3-aprobaciones/',
        data={'no_cia': '01', 'archivo': _as_upload(archivo)},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert len(body['resultados']) == 1
    assert body['resultados'][0]['ok'] is True
    assert body['resultados'][0]['estado'] == 'Aprobacion Comercial Aprobada.'
    assert envios == [('certecf', 'E310000000001', '130217432')]


def test_paso4_factura_real_requiere_login(client, db):
    resp = client.post('/api/fe/certificacion/paso4-factura-real/', data={'no_cia': '01'})
    assert resp.status_code in (302, 401, 403)


def test_paso4_factura_real_campos_requeridos(cliente_autenticado):
    resp = cliente_autenticado.post(
        '/api/fe/certificacion/paso4-factura-real/',
        data={'no_cia': '01', 'tipo_ecf': '31'})
    assert resp.status_code == 400
    assert 'punto' in resp.json()['detail'].lower()


def test_paso4_factura_real_envia_tipo_31(cliente_autenticado, monkeypatch):
    monkeypatch.setattr(
        ecf_builder, 'construir_ecf_31',
        lambda no_cia, punto, tipo_factura, no_factura: '<ECF><Encabezado><IdDoc><eNCF>E310000000054</eNCF></IdDoc></Encabezado></ECF>')
    envios = []

    def fake_enviar_ecf(no_cia, ambiente, e_ncf, xml_sin_firmar):
        envios.append((ambiente, e_ncf))
        return {'trackId': 'TRACK-1', 'xml_firmado': '<ECF firmado/>',
                'respuesta_cruda': {'trackId': 'TRACK-1'}}

    monkeypatch.setattr(dgii_client, 'enviar_ecf', fake_enviar_ecf)
    monkeypatch.setattr(
        'apps.legacy.repositories.fe_repo.save_documento_enviado',
        lambda *a, **k: None)

    resp = cliente_autenticado.post(
        '/api/fe/certificacion/paso4-factura-real/',
        data={'no_cia': '01', 'tipo_ecf': '31', 'punto': '01',
              'tipo_factura': 'FT', 'no_factura': '123'})
    assert resp.status_code == 200
    body = resp.json()
    assert body['ok'] is True
    assert body['trackId'] == 'TRACK-1'
    assert body['encf'] == 'E310000000054'
    assert envios == [('certecf', 'E310000000054')]


def test_paso4_factura_real_tipo_invalido_da_400(cliente_autenticado):
    resp = cliente_autenticado.post(
        '/api/fe/certificacion/paso4-factura-real/',
        data={'no_cia': '01', 'tipo_ecf': '99', 'punto': '01',
              'tipo_factura': 'FT', 'no_factura': '123'})
    assert resp.status_code == 400
    assert '31' in resp.json()['detail'] or '32' in resp.json()['detail']


def test_paso4_factura_real_error_de_builder_da_400(cliente_autenticado, monkeypatch):
    def fake_construir(*a, **k):
        raise ecf_builder.ECFBuilderError('factura anulada')

    monkeypatch.setattr(ecf_builder, 'construir_ecf_32', fake_construir)

    resp = cliente_autenticado.post(
        '/api/fe/certificacion/paso4-factura-real/',
        data={'no_cia': '01', 'tipo_ecf': '32', 'punto': '01',
              'tipo_factura': 'FC', 'no_factura': '456'})
    assert resp.status_code == 400
    assert 'factura anulada' in resp.json()['detail']
